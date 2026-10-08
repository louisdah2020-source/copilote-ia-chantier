"""
reports/excel_exporter.py
Ré-exportateur de classeur Excel enrichi.
Prend le classeur d'origine et réécrit les feuilles AVANCEMENT et ALERTES
avec les calculs et détections à jour.
"""
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from core.models import KPISummary, AlertItem
from core.excel_generator import HEADER_FILL, HEADER_FONT, REGULAR_FONT, THIN_BORDER, _apply_table_styling


class ExcelExporter:
    """Met à jour et exporte le classeur Excel avec les feuilles AVANCEMENT et ALERTES régénérées."""

    def __init__(self, source_path: str, kpi: KPISummary, alerts: list[AlertItem]):
        self.source_path = Path(source_path)
        self.kpi = kpi
        self.alerts = alerts

    def export(self, output_path: str) -> str:
        """Génère le fichier Excel enrichi."""
        wb = openpyxl.load_workbook(self.source_path)

        # 1. Mise à jour de la feuille AVANCEMENT
        if "AVANCEMENT" in wb.sheetnames:
            ws_av = wb["AVANCEMENT"]
            # Conserver l'en-tête, vider les données
            while ws_av.max_row > 1:
                ws_av.delete_rows(2)
        else:
            ws_av = wb.create_sheet(title="AVANCEMENT")
            ws_av.append(["Date", "Lot", "% prévu", "% réalisé", "Écart", "Tendance"])

        for lot_data in self.kpi.repartition_lots:
            ecart = lot_data["ecart_avancement"]
            if ecart < -5:
                tendance = "🔴"
            elif ecart < 0:
                tendance = "🟠"
            else:
                tendance = "🟢"

            ws_av.append([
                self.kpi.date_analyse,
                lot_data["lot"],
                lot_data["pct_prevu"] / 100.0,
                lot_data["pct_reel"] / 100.0,
                ecart / 100.0,
                tendance
            ])
        _apply_table_styling(ws_av)

        # 2. Mise à jour de la feuille ALERTES
        if "ALERTES" in wb.sheetnames:
            ws_alt = wb["ALERTES"]
            while ws_alt.max_row > 1:
                ws_alt.delete_rows(2)
        else:
            ws_alt = wb.create_sheet(title="ALERTES")
            ws_alt.append(["Date", "Type", "Lot", "Gravité", "Description", "Action recommandée"])

        for alt in self.alerts:
            ws_alt.append([
                alt.date,
                alt.type_alerte,
                alt.lot,
                alt.gravite,
                alt.description,
                alt.action_recommandee
            ])
        _apply_table_styling(ws_alt)

        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        return output_path
